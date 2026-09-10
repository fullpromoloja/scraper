import asyncio
import getpass
import json
import os
import re
import sys
from dataclasses import dataclass
from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from playwright.async_api import async_playwright, TimeoutError as PlaywrightTimeoutError

# ========================= CONFIGURAÇÕES =========================
EMPRESA = "IMPORTEC_ATACADO"  # Ajustado de IMPORTE para IMPORTEC
LUCRO_DESEJADO = 0.15  # 15% de lucro sobre o custo
URL_PADRAO = "https://www.importecatacado.com.br/categoria/brinquedos"
PAGE_PARAM = "pagina"
MAX_PAGINAS = 200
COOKIES_FILE = "cookies.json"

# Pega a URL informada na linha de comando ou usa a padrão
URL_CATEGORIA = sys.argv[1] if len(sys.argv) > 1 else URL_PADRAO

# ========================= FUNÇÃO DE VALIDAÇÃO DE COOKIES =========================
def validar_e_limpar_cookies(caminho_arquivo: str) -> bool:
    """Verifica se o arquivo de cookies existe e contém um JSON válido."""
    if not os.path.exists(caminho_arquivo):
        return False

    try:
        with open(caminho_arquivo, "r", encoding="utf-8") as f:
            dados = json.load(f)
            if isinstance(dados, dict) and ("cookies" in dados or "origins" in dados):
                return True
    except Exception:
        pass

    print("⚠️  Arquivo 'cookies.json' inválido ou corrompido. Limpando arquivo...")
    try:
        os.remove(caminho_arquivo)
    except OSError:
        pass
    return False

# ========================= MODELO DE DADOS =========================
@dataclass
class Produto:
    nome: str
    custo: float
    lucro_desejado: float
    liquido_necessario: float
    preco_venda_necessario: float

# ========================= REGRAS DE TAXA SHOPEE =========================
FAIXAS_SHOPEE = [
    (79.99, 0.20, 4.0),
    (99.99, 0.14, 16.0),
    (199.99, 0.14, 20.0),
    (float("inf"), 0.14, 26.0),
]

def calcular_preco_venda_necessario(liquido_necessario: float) -> float:
    limite_inferior = 0.0

    for limite_superior, percentual, fixo in FAIXAS_SHOPEE:
        preco_candidato = (liquido_necessario + fixo) / (1 - percentual)
        dentro_da_faixa = limite_inferior <= preco_candidato <= limite_superior

        if limite_inferior == 0.0:
            dentro_da_faixa = preco_candidato <= limite_superior

        if dentro_da_faixa:
            return round(preco_candidato, 2)

        limite_inferior = limite_superior

    percentual, fixo = FAIXAS_SHOPEE[-1][1], FAIXAS_SHOPEE[-1][2]
    return round((liquido_necessario + fixo) / (1 - percentual), 2)

def calcular_precificacao(custo: float, lucro_desejado: float = LUCRO_DESEJADO) -> Produto:
    liquido_necessario = round(custo * (1 + lucro_desejado), 2)
    preco_venda = calcular_preco_venda_necessario(liquido_necessario)

    return Produto(
        nome="",
        custo=custo,
        lucro_desejado=lucro_desejado,
        liquido_necessario=liquido_necessario,
        preco_venda_necessario=preco_venda,
    )

def parse_preco(texto: str) -> Optional[float]:
    if not texto:
        return None
    match = re.search(r"R\$\s*([\d.,]+)", texto)
    if not match:
        return None
    limpo = match.group(1).replace(".", "").replace(",", ".")
    try:
        return float(limpo)
    except ValueError:
        return None

# ========================= EXCEL =========================
def gerar_planilha(produtos: list, caminho_saida: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Precificacao"

    cabecalho = [
        "Nome Produto",
        "Custo Produto",
        "Lucro Desejado",
        "Líquido Necessário",
        "Preço Necessário para Venda",
    ]

    ws.append(cabecalho)

    for col in range(1, len(cabecalho) + 1):
        ws.cell(row=1, column=col).font = Font(bold=True)

    for nome, custo in produtos:
        resultado = calcular_precificacao(custo)
        ws.append([
            nome,
            resultado.custo,
            resultado.lucro_desejado,
            resultado.liquido_necessario,
            resultado.preco_venda_necessario,
        ])

    ws.column_dimensions[get_column_letter(1)].width = 50
    for col in range(2, 6):
        ws.column_dimensions[get_column_letter(col)].width = 22

    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = 'R$ #,##0.00'

    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.number_format = "0%"

    for row in ws.iter_rows(min_row=2, min_col=4, max_col=5):
        for cell in row:
            cell.number_format = 'R$ #,##0.00'

    wb.save(caminho_saida)
    print(f"\n✨ Planilha gerada com sucesso: {caminho_saida}")

# ========================= EXECUÇÃO PLAYWRIGHT =========================
async def run():
    usuario_input = ""
    senha_input = ""

    tem_cookies_validos = validar_e_limpar_cookies(COOKIES_FILE)

    if not tem_cookies_validos:
        print("\n" + "┌" + "─" * 58 + "┐")
        print("│" + "AUTENTICAÇÃO - IMPORTEC ATACADO".center(58) + "│")
        print("└" + "─" * 58 + "┘")
        usuario_input = input("  📧 E-mail / Usuário : ").strip()
        senha_input = getpass.getpass("  🔑 Senha            : ").strip()

        if not usuario_input or not senha_input:
            print("\n❌ E-mail e senha são obrigatórios.")
            sys.exit(1)
    else:
        print(f"🔒 Sessão encontrada em '{COOKIES_FILE}'. Carregando cookies...")

    print("\n" + "┌" + "─" * 58 + "┐")
    print("│" + "INICIANDO SCRAPING".center(58) + "│")
    print("└" + "─" * 58 + "┘")

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=False)
        
        if validar_e_limpar_cookies(COOKIES_FILE):
            context = await browser.new_context(storage_state=COOKIES_FILE)
        else:
            context = await browser.new_context()

        page = await context.new_page()

        try:
            print("🌐 Acessando a página inicial...")
            await page.goto("https://www.importecatacado.com.br/", wait_until="networkidle")

            if not await page.is_visible(".f1-client-info--logged, .js-client-logout"):
                if not usuario_input or not senha_input:
                    print("\n⚠️  Sessão expirada. Por favor, faça login novamente.")
                    usuario_input = input("  📧 E-mail / Usuário : ").strip()
                    senha_input = getpass.getpass("  🔑 Senha            : ").strip()

                if await page.is_visible(".js-client-login"):
                    print("🔄 Abrindo modal de login...")
                    await page.click(".js-client-login")
                    await page.wait_for_selector("#formLogin", state="visible", timeout=5000)

                print("✍️  Preenchendo credenciais...")
                await page.fill("#formLogin #username", usuario_input)
                await page.fill("#formLogin #password", senha_input)

                print("🚀 Enviando dados de acesso...")
                async with page.expect_response("**/cliente/entrar**"):
                    await page.click(".f1-modal-login__submit")

                await page.wait_for_selector(".f1-client-info--logged, .js-client-logout", timeout=12000)
                print("✅ Login efetuado com sucesso!")

                await context.storage_state(path=COOKIES_FILE)
                print(f"💾 Sessão salva em '{COOKIES_FILE}'.")
            else:
                print("✅ Login confirmado via cookies arquivados!")

        except PlaywrightTimeoutError:
            print("\n❌ [ERRO DE AUTENTICAÇÃO]: Falha ao realizar o login.")
            if os.path.exists(COOKIES_FILE):
                os.remove(COOKIES_FILE)
            await browser.close()
            sys.exit(1)
        except Exception as e:
            print(f"\n❌ [ERRO INESPERADO]: {e}")
            await browser.close()
            sys.exit(1)

        produtos_extraidos = []

        # Loop de Paginação
        for pagina in range(1, MAX_PAGINAS + 1):
            separador = "&" if "?" in URL_CATEGORIA else "?"
            url_pagina = f"{URL_CATEGORIA}{separador}{PAGE_PARAM}={pagina}"

            print(f"\n🔍 [Página {pagina}] Acessando: {url_pagina}")
            
            try:
                await page.goto(url_pagina, wait_until="networkidle")
                await page.wait_for_selector(".f1-product-item", timeout=10000)
            except PlaywrightTimeoutError:
                print(f"🛑 Nenhum produto encontrado na página {pagina}. Finalizando paginação.")
                break
            except Exception as e:
                print(f"❌ Erro ao carregar a página {pagina}: {e}")
                break

            produtos_dom = await page.query_selector_all(".f1-product-item")
            print(f"📦 Elementos encontrados: {len(produtos_dom)}")

            if not produtos_dom:
                print("🛑 Fim dos produtos na categoria.")
                break

            produtos_pagina = 0
            for produto in produtos_dom:
                nome_elem = await produto.query_selector(".f1-product-item__name-link")
                preco_elem = await produto.query_selector(".f1-box-price__price")

                nome_texto = await nome_elem.inner_text() if nome_elem else ""
                preco_texto = await preco_elem.inner_text() if preco_elem else ""

                custo = parse_preco(preco_texto)

                if nome_texto and custo is not None:
                    produtos_extraidos.append((nome_texto.strip(), custo))
                    produtos_pagina += 1

            if produtos_pagina == 0:
                print("🛑 Nenhum produto válido extraído nesta página. Finalizando.")
                break

        await browser.close()

        # Deduplicação
        vistos = set()
        produtos_unicos = []
        for item in produtos_extraidos:
            if item not in vistos:
                vistos.add(item)
                produtos_unicos.append(item)

        print(f"\n📊 Total de {len(produtos_unicos)} produtos únicos coletados.")

        if produtos_unicos:
            categoria_nome = URL_CATEGORIA.rstrip("/").split("/")[-1].upper()
            arquivo_saida = f"{EMPRESA}_{categoria_nome}.xlsx"
            gerar_planilha(produtos_unicos, arquivo_saida)
        else:
            print("⚠️  Nenhum produto com preço válido foi encontrado para exportação.")

if __name__ == "__main__":
    asyncio.run(run())