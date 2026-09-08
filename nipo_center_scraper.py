#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Raspagem de produtos do Nipo Center (atacadonipocenter.com.br) + cálculo
do preço de venda necessário no Shopee para atingir a margem de lucro
desejada sobre o custo de compra (valor bruto).

COMO USAR
---------
1. Login — duas opções, escolha uma:

   OPÇÃO A (recomendada, mais simples): cookies do navegador
   -----------------------------------------------------------
   a) Faça login no site normalmente pelo Chrome (com o CNPJ cadastrado).
   b) Exporte os cookies da sessão logada para um arquivo `cookies.json`
      (extensão "Cookie-Editor" -> no site logado -> Export -> Export as JSON).
      Coloque o arquivo na mesma pasta deste script.

   OPÇÃO B: login automático via email/senha
   ------------------------------------------
   Só funciona se LOGIN_URL e os nomes dos campos abaixo (CONFIG DE LOGIN)
   estiverem corretos — capture isso no DevTools do Chrome (F12 -> aba
   Network -> faça login -> ache a requisição POST -> Copy as cURL) e ajuste
   a função fazer_login() mais abaixo de acordo.

   Defina as variáveis de ambiente antes de rodar:
     export NIPO_EMAIL="seu-email@exemplo.com"
     export NIPO_SENHA="sua-senha"

   (Nunca deixe a senha escrita diretamente no código.)

2. Abra uma página de categoria já logado, aperte F12 (DevTools), inspecione
   um card de produto e ajuste os seletores na seção CONFIG abaixo:
   - PRODUCT_CARD_SELECTOR
   - PRODUCT_NAME_SELECTOR
   - PRODUCT_PRICE_SELECTOR

3. Rode passing a(s) URL(s) de categoria como argumento:

   python nipo_center_scraper.py https://www.atacadonipocenter.com.br/categoria/brinquedos

   (pode passar mais de uma URL separada por espaço)

   Exemplo:
   python nipo_center_scraper.py \
       https://www.atacadonipocenter.com.br/categoria/brinquedos \
       https://www.atacadonipocenter.com.br/categoria/ferramentas

   O arquivo será gerado automaticamente no formato:

   NIPOCENTER_BRINQUEDOS.xlsx

   Ou, caso sejam informadas várias categorias:

   NIPOCENTER_BRINQUEDOS_FERRAMENTAS.xlsx
"""

import argparse
import json
import os
import re
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


# ========================= CONFIG =========================

# Nome da empresa/fornecedor usado no nome do arquivo Excel.
# Exemplo:
# NIPOCENTER + BRINQUEDOS -> NIPOCENTER_BRINQUEDOS.xlsx
EMPRESA = "NIPOCENTER"

BASE_URL = "https://www.atacadonipocenter.com.br"

COOKIES_FILE = "cookies.json"

# --- CONFIG DE LOGIN (opção B, só usada se cookies.json não existir) ---
LOGIN_URL = "https://www.atacadonipocenter.com.br/cliente/entrar"
LOGIN_FIELD_EMAIL = "email"
LOGIN_FIELD_SENHA = "senha"

LUCRO_DESEJADO = 0.15  # 15%

# Parâmetro de paginação observado no site: ?pagina=N
PAGE_PARAM = "pagina"
MAX_PAGES_PER_CATEGORY = 200  # trava de segurança
REQUEST_DELAY_SECONDS = 1.5   # educado com o servidor

# --- SELETORES CSS (confirmados a partir do HTML real de um card) ---
PRODUCT_CARD_SELECTOR = "div.product-information"
PRODUCT_NAME_SELECTOR = "a.b2b-product-link"
PRODUCT_PRICE_SELECTOR = "span.new-price"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0 Safari/537.36"
    )
}


# ========================= MODELO DE DADOS =========================


@dataclass
class Produto:
    nome: str
    custo: float
    lucro_desejado: float
    liquido_necessario: float
    preco_venda_necessario: float


# ========================= REGRAS DE TAXA SHOPEE =========================

FAIXAS_SHOPEE = [
    # (limite_superior_da_faixa, percentual, valor_fixo)
    (79.99, 0.20, 4.0),
    (99.99, 0.14, 16.0),
    (199.99, 0.14, 20.0),
    (float("inf"), 0.14, 26.0),
]


def calcular_preco_venda_necessario(liquido_necessario: float) -> float:
    limite_inferior = 0.0

    for limite_superior, percentual, fixo in FAIXAS_SHOPEE:
        preco_candidato = (liquido_necessario + fixo) / (1 - percentual)

        dentro_da_faixa = limite_inferior <= preco_candidato <= limite_superior

        if limite_inferior == 0.0:
            dentro_da_faixa = preco_candidato <= limite_superior

        if dentro_da_faixa:
            return round(preco_candidato, 2)

        limite_inferior = limite_superior

    percentual, fixo = FAIXAS_SHOPEE[-1][1], FAIXAS_SHOPEE[-1][2]
    return round((liquido_necessario + fixo) / (1 - percentual), 2)


def calcular_precificacao(
    custo: float,
    lucro_desejado: float = LUCRO_DESEJADO
) -> Produto:

    liquido_necessario = round(custo * (1 + lucro_desejado), 2)
    preco_venda = calcular_preco_venda_necessario(liquido_necessario)

    return Produto(
        nome="",
        custo=custo,
        lucro_desejado=lucro_desejado,
        liquido_necessario=liquido_necessario,
        preco_venda_necessario=preco_venda,
    )


# ========================= RASPAGEM =========================


def carregar_cookies(session: requests.Session, path: str) -> bool:
    cookie_path = Path(path)

    if not cookie_path.exists():
        return False

    try:
        with open(cookie_path, "r", encoding="utf-8") as f:
            cookies = json.load(f)
    except (OSError, json.JSONDecodeError) as e:
        print(f"[ERRO] Não foi possível ler {path}: {e}")
        return False

    for cookie in cookies:
        session.cookies.set(
            cookie.get("name"),
            cookie.get("value"),
            domain=cookie.get("domain")
        )

    print(f"[OK] {len(cookies)} cookies carregados de {path}.")
    return True


def fazer_login(session: requests.Session) -> bool:
    email = os.environ.get("NIPO_EMAIL")
    senha = os.environ.get("NIPO_SENHA")

    if not email or not senha:
        return False

    payload = {
        LOGIN_FIELD_EMAIL: email,
        LOGIN_FIELD_SENHA: senha,
    }

    try:
        resp = session.post(
            LOGIN_URL,
            data=payload,
            headers=HEADERS,
            timeout=20
        )
        resp.raise_for_status()
    except requests.RequestException as e:
        print(f"[ERRO] Falha ao tentar logar: {e}")
        return False

    print("[OK] Requisição de login enviada (confira se autenticou de fato).")
    return True


def autenticar(session: requests.Session) -> None:
    if carregar_cookies(session, COOKIES_FILE):
        return

    if fazer_login(session):
        return

    print(
        "[AVISO] Nenhuma forma de login configurada "
        "(nem cookies.json, nem NIPO_EMAIL/NIPO_SENHA). "
        "Continuando sem login — os preços provavelmente não vão aparecer."
    )


def parse_preco(texto: str) -> Optional[float]:
    if not texto:
        return None

    match = re.search(r"R\$\s*([\d.,]+)", texto)
    if not match:
        return None

    limpo = match.group(1).replace(".", "").replace(",", ".")

    try:
        return float(limpo)
    except ValueError:
        return None


def fetch_page(session: requests.Session, url: str) -> str:
    resp = session.get(url, headers=HEADERS, timeout=20)
    resp.raise_for_status()
    return resp.text


def extrair_produtos_da_pagina(html: str) -> list:
    soup = BeautifulSoup(html, "html.parser")
    cards = soup.select(PRODUCT_CARD_SELECTOR)
    produtos = []

    for card in cards:
        nome_el = card.select_one(PRODUCT_NAME_SELECTOR)
        preco_el = card.select_one(PRODUCT_PRICE_SELECTOR)

        if not nome_el or not preco_el:
            continue

        nome = nome_el.get_text(strip=True)
        preco = parse_preco(preco_el.get_text(strip=True))

        if nome and preco:
            produtos.append((nome, preco))

    return produtos


def raspar_categoria(session: requests.Session, categoria_url: str) -> list:
    produtos_categoria = []

    for pagina in range(1, MAX_PAGES_PER_CATEGORY + 1):
        separador = "&" if "?" in categoria_url else "?"
        url = f"{categoria_url}{separador}{PAGE_PARAM}={pagina}"

        print(f"  -> Página {pagina}: {url}")

        try:
            html = fetch_page(session, url)
        except requests.RequestException as e:
            print(f"     [ERRO] Falha ao buscar {url}: {e}")
            break

        produtos = extrair_produtos_da_pagina(html)

        if not produtos:
            print("     Nenhum produto encontrado — fim da categoria.")
            break

        produtos_categoria.extend(produtos)
        time.sleep(REQUEST_DELAY_SECONDS)

    return produtos_categoria


# ========================= EXCEL =========================


def gerar_planilha(produtos: list, caminho_saida: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Precificacao"

    cabecalho = [
        "Nome Produto",
        "Custo Produto",
        "Lucro Desejado",
        "Líquido Necessário",
        "Preço Necessário para Venda",
    ]

    ws.append(cabecalho)

    # Cabeçalho em negrito
    for col in range(1, len(cabecalho) + 1):
        ws.cell(row=1, column=col).font = Font(bold=True)

    # Inserção dos produtos
    for nome, custo in produtos:
        resultado = calcular_precificacao(custo)
        ws.append([
            nome,
            resultado.custo,
            resultado.lucro_desejado,
            resultado.liquido_necessario,
            resultado.preco_venda_necessario,
        ])

    # Formatação das colunas
    ws.column_dimensions[get_column_letter(1)].width = 50
    for col in range(2, 6):
        ws.column_dimensions[get_column_letter(col)].width = 20

    # Formato monetário - Custo
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = 'R$ #,##0.00'

    # Formato percentual
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.number_format = "0%"

    # Formato monetário
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=5):
        for cell in row:
            cell.number_format = 'R$ #,##0.00'

    wb.save(caminho_saida)
    print(f"\n[OK] Planilha salva em: {caminho_saida}")


# ========================= ARGUMENTOS =========================


def parse_argumentos() -> list:
    parser = argparse.ArgumentParser(
        description="Raspa produtos do Nipo Center por categoria e gera planilha de precificação."
    )
    parser.add_argument(
        "categorias",
        nargs="+",
        help="Uma ou mais URLs de categoria."
    )

    args = parser.parse_args()

    categorias_validas = [
        url for url in args.categorias
        if url.startswith("http://") or url.startswith("https://")
    ]

    if not categorias_validas:
        parser.error("Nenhuma URL válida de categoria foi informada.")

    return categorias_validas


def nome_categoria_da_url(url: str) -> str:
    """
    Extrai a categoria da URL e limpa o formato para ser usado no nome do arquivo.
    Exemplo:
        https://.../categoria/brinquedos -> BRINQUEDOS
        https://.../categoria/brinquedos-e-jogos -> BRINQUEDOS_E_JOGOS
    """
    caminho = url.rstrip("/").split("?")[0]
    categoria = caminho.split("/")[-1]
    
    # Substitui hífens e caracteres especiais por underline para um nome limpo
    categoria_limpa = re.sub(r"[^a-zA-Z0-9]", "_", categoria)
    
    return categoria_limpa.upper()


# ========================= MAIN =========================


def main():
    category_urls = parse_argumentos()

    session = requests.Session()
    autenticar(session)

    todos_produtos = []

    for categoria_url in category_urls:
        print(f"\nRaspando categoria: {categoria_url}")
        produtos = raspar_categoria(session, categoria_url)
        print(f"  {len(produtos)} produtos encontrados nesta categoria.")
        todos_produtos.extend(produtos)

    # Extrai o nome de cada categoria e junta com '_'
    nomes_categorias = "_".join(
        nome_categoria_da_url(url) for url in category_urls
    )

    # Gera o nome no formato desejado: EMPRESA_CATEGORIA.xlsx
    output_xlsx = f"{EMPRESA}_{nomes_categorias}.xlsx"

    print(f"\n[INFO] Arquivo de saída: {output_xlsx}")

    if not todos_produtos:
        print(
            "\n[ATENÇÃO] Nenhum produto foi extraído.\n"
            "Verifique os seletores CSS, estado do login ou se o site requer Playwright."
        )
        return

    # Remove duplicados mantendo a ordem
    vistos = set()
    produtos_unicos = []

    for nome, preco in todos_produtos:
        chave = (nome, preco)
        if chave not in vistos:
            vistos.add(chave)
            produtos_unicos.append((nome, preco))

    print(f"[INFO] {len(produtos_unicos)} produtos únicos encontrados.")

    # Salva com o nome dinâmico
    gerar_planilha(produtos_unicos, output_xlsx)


if __name__ == "__main__":
    main()